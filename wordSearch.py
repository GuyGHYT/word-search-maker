import random
import string
from docx import Document
from docx.shared import Pt

def create_word_search(words, grid_size=20):
    grid = [[' ' for _ in range(grid_size)] for _ in range(grid_size)]
    
    def place_word(word):
        word_len = len(word)
        placed = False
        attempts = 0
        while not placed and attempts < 100:
            direction = random.choice(['horizontal', 'vertical', 'diagonal'])
            if direction == 'horizontal':
                row = random.randint(0, grid_size - 1)
                col = random.randint(0, grid_size - word_len)
                if all(grid[row][col + i] in (' ', word[i]) for i in range(word_len)):
                    for i in range(word_len):
                        grid[row][col + i] = word[i]
                    placed = True
            elif direction == 'vertical':
                row = random.randint(0, grid_size - word_len)
                col = random.randint(0, grid_size - 1)
                if all(grid[row + i][col] in (' ', word[i]) for i in range(word_len)):
                    for i in range(word_len):
                        grid[row + i][col] = word[i]
                    placed = True
            elif direction == 'diagonal':
                row = random.randint(0, grid_size - word_len)
                col = random.randint(0, grid_size - word_len)
                if all(grid[row + i][col + i] in (' ', word[i]) for i in range(word_len)):
                    for i in range(word_len):
                        grid[row + i][col + i] = word[i]
                    placed = True
            attempts += 1
    
    for word in words:
        place_word(word.upper())
    
    for row in range(grid_size):
        for col in range(grid_size):
            if grid[row][col] == ' ':
                grid[row][col] = random.choice(string.ascii_uppercase)
    
    return grid

def save_grid_to_word(grid, filename):
    doc = Document()
    table = doc.add_table(rows=len(grid), cols=len(grid[0]))
    table.style = 'Table Grid'
    
    for i, row in enumerate(grid):
        for j, cell in enumerate(row):
            table.cell(i, j).text = cell
            table.cell(i, j).paragraphs[0].runs[0].font.size = Pt(12)
            table.cell(i, j).paragraphs[0].alignment = 1  # Center alignment
    
    doc.save(filename)

words = [
    "Americans", "Wall", "Republican", "President", "GulfofAmerica",
    "Conservative", "Determined", "Guns", "Twitter", "DOGE", "Musk",
    "Pardon", "Veto", "Cabinet", "FBI", "Amendment", "Tarriff",
    "Aluminium", "Steel", "Divisive"
]
grid = create_word_search(words, grid_size=20)
save_grid_to_word(grid, 'word_search.docx')
